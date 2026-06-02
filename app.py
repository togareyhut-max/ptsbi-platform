import csv
import io
import json
import os
from datetime import datetime
from pathlib import Path

try:
    from dotenv import load_dotenv
except ImportError:
    def load_dotenv(*_args, **_kwargs):
        return False

_APP_DIR = Path(__file__).resolve().parent
load_dotenv(_APP_DIR / ".env")
# Dev lokal: kosongkan DATABASE_URL atau set USE_SQLITE=1 di .env
if not os.environ.get("DATABASE_URL", "").strip():
    os.environ.setdefault("USE_SQLITE", "1")

from flask import Flask, flash, g, has_request_context, jsonify, redirect, render_template, request, send_file, session, url_for
from werkzeug.utils import secure_filename
from werkzeug.security import check_password_hash, generate_password_hash

from db import (
    connect,
    coerce_submitted_by,
    init_schema,
    insert_returning_id,
    migrate_schema,
    reset_sqlite_database_file,
    sql_approved,
    sql_not_approved,
    sqlite_schema_ok,
)
from services.auth_roles import (
    developer_required,
    is_staff_role,
    login_required,
    role_label,
    staff_required,
    staff_redirect_endpoint,
    user_can_view_tarombo,
)
from services.audit import log_audit
from services.form_config import (
    get_all_fields_admin,
    get_fields_by_section,
    seed_form_fields,
    validate_submission_from_config,
)
from services.batak_names import format_person_display_name, format_spouse_display_name
from services.matching import find_existing_person, normalize_name, person_identity_fields
from services.graph_sync import place_person_in_tree, reconcile_tree, sync_tarombo
from services.person_merge import backfill_mother_name_normalized, find_duplicate_identity_groups
from services.tree_display import collect_tree_person_ids
from services.opung import (
    build_relationship_maps,
    enrich_people_panggoaran_for_tree,
    format_panggoaran_display,
    get_opung_hover_suffix,
    recalculate_all_panggoaran,
    sort_person_ids,
)
from services.partuturan import PARTUTURAN_STATUSES, partuturan_detail, partuturan_label
from services.sundut import (
    effective_sundut,
    list_sundut_entries,
    seed_sundut_entries,
    sundut_title_map,
)
from services.db_admin import (
    BACKUP_DIR,
    create_backup,
    list_backups,
    reset_database_file_sqlite,
    resolve_backup_path,
    restore_backup,
    restore_uploaded_file,
    wipe_all_data,
)
from services.member_registration import create_member_with_tarombo
from services.wp_integration_api import register_wp_api
from services.tree_editor import apply_tree_changes
from services.sibling_order import (
    apply_child_orders,
    assign_orders_from_birth_year,
    get_siblings_for_person,
    order_label,
    parse_child_order,
    sort_people_list,
)
from services.tarombo_person import add_tarombo_person, parse_tarombo_person_form
from services.tree_display import build_unique_tree
from services.site_branding import (
    custom_logo_path,
    ensure_upload_dir,
    logo_info,
    remove_custom_logo,
    save_logo_upload,
    site_logo_url,
)
from services.membership_api import bp as membership_api_bp, upsert_member_profile_record

app = Flask(__name__)
app.register_blueprint(membership_api_bp)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "change-this-in-production")
register_wp_api(app)
_db_bootstrapped = False


def get_db():
    if "db" not in g:
        g.db = connect()
    return g.db


def reset_request_db():
    if not has_request_context():
        return
    db = g.pop("db", None)
    if db is not None:
        try:
            db._conn.close()
        except Exception:
            pass


def close_database_for_restore():
    """Tutup koneksi request sebelum timpa file SQLite / restore."""
    global _db_bootstrapped
    reset_request_db()
    _db_bootstrapped = False


@app.teardown_appcontext
def close_db(_error):
    db = g.pop("db", None)
    if db is not None:
        db._conn.close()


def backfill_normalized_names(db):
    try:
        rows = db.fetchall(
            """
            SELECT id, full_name, marga, father_name, mother_name, birth_year
            FROM people
            WHERE name_normalized IS NULL OR name_normalized = ''
               OR mother_name_normalized IS NULL
            """
        )
    except Exception:
        return
    for row in rows:
        ident = person_identity_fields(
            row["full_name"],
            row["marga"],
            row["father_name"],
            row["birth_year"],
            row.get("mother_name"),
        )
        db.execute(
            """
            UPDATE people SET
              name_normalized = ?,
              father_name_normalized = ?,
              mother_name_normalized = ?
            WHERE id = ?
            """,
            (
                ident["name_normalized"],
                ident["father_name_normalized"],
                ident["mother_name_normalized"],
                row["id"],
            ),
        )
    if rows:
        db.commit()
    try:
        backfill_mother_name_normalized(db)
    except Exception:
        app.logger.exception("backfill_mother_name_normalized gagal")


def init_db():
    db = get_db()
    if db.backend == "sqlite" and not sqlite_schema_ok(db):
        try:
            db._conn.close()
        except Exception:
            pass
        reset_request_db()
        reset_sqlite_database_file()
        db = get_db()

    try:
        init_schema(db)
    except Exception:
        app.logger.exception("init_schema gagal")
        if db.backend == "sqlite":
            try:
                db._conn.close()
            except Exception:
                pass
            reset_request_db()
            reset_sqlite_database_file()
            db = get_db()
            init_schema(db)

    migrate_schema(db)
    ensure_upload_dir()
    backfill_normalized_names(db)
    try:
        seed_form_fields(db)
    except Exception:
        app.logger.exception("seed_form_fields gagal")
    try:
        seed_sundut_entries(db)
    except Exception:
        app.logger.exception("seed_sundut_entries gagal")


def ensure_public_tree_data(db):
    """Pastikan ada data approved agar pohon tarombo tidak kosong setelah setup pertama."""
    row = db.fetchone(f"SELECT COUNT(*) AS c FROM people WHERE {sql_approved(db)}")
    if row and int(row["c"]) > 0:
        return
    seed_demo_tree(db)


def bootstrap_database(*, force=False):
    global _db_bootstrapped
    if force:
        _db_bootstrapped = False
        reset_request_db()
    if _db_bootstrapped:
        return
    try:
        init_db()
        seed_data()
        _db_bootstrapped = True
    except Exception:
        app.logger.exception("bootstrap_database gagal")
        if force:
            raise


@app.context_processor
def inject_site_branding():
    try:
        return {
            "site_logo_url": site_logo_url(),
            "can_view_tarombo": user_can_view_tarombo(),
            "is_staff": is_staff_role(),
            "is_developer": session.get("role") == "developer",
            "role_label": role_label(session.get("role")),
        }
    except Exception:
        return {
            "site_logo_url": "/static/logo-ptsbi.svg",
            "can_view_tarombo": False,
            "is_staff": False,
            "is_developer": False,
            "role_label": "Member",
        }


@app.route("/site-logo.png")
def site_logo():
    path = custom_logo_path()
    if not path:
        return redirect(url_for("static", filename="logo-ptsbi.svg"))
    return send_file(path, mimetype="image/png", max_age=86400)


@app.before_request
def _ensure_db_ready():
    if request.endpoint in (None, "static"):
        return
    if not _db_bootstrapped:
        try:
            bootstrap_database()
        except Exception:
            app.logger.exception("Database bootstrap gagal pada request")


def register_form_context(db):
    sections, titles = get_fields_by_section(db)
    return {"form_sections": sections, "section_titles": titles}


def admin_tarombo_form_context(db):
    """Form admin: hanya data tarombo, tanpa akun login."""
    sections, titles = get_fields_by_section(db)
    tarombo_only = {k: v for k, v in sections.items() if k in ("person", "tarombo", "marriage")}
    return {"form_sections": tarombo_only, "section_titles": titles}


def add_tarombo_person_from_request(db, form, *, admin_user_id):
    data = parse_tarombo_person_form(form)
    error = validate_tarombo_form(
        gender=data["gender"],
        marga=data["marga"],
        tarombo_status=data["tarombo_status"],
        father_name=data["father_name"],
        mother_name=data["mother_name"],
        reference_female_line_name=data["reference_female_line_name"],
        is_married=data["is_married"],
        spouse_name=data["spouse_name"],
        spouse_marga=data["spouse_marga"],
    )
    if error:
        raise ValueError(error)
    result = add_tarombo_person(
        db,
        data,
        bool_db=bool_db,
        admin_user_id=coerce_submitted_by(db, admin_user_id),
        person_identity_fields=person_identity_fields,
        normalize_spouse_fields=normalize_spouse_fields,
        save_submitted_children=save_submitted_children,
        find_existing_person=find_existing_person,
    )
    finalize_new_tarombo_person(db, result["person_id"])
    return result


def bool_db(db, value: bool):
    if db.backend == "postgres":
        return value
    return 1 if value else 0


def approve_member_registration(db, user_id: int, *, actor_user_id=None):
    """Setujui pendaftaran member: aktifkan akun + data tarombo + tempatkan di pohon."""
    user_id = int(user_id)
    db.execute(
        "UPDATE users SET member_status = 'active' WHERE id = ? AND role = 'member'",
        (user_id,),
    )
    rows = db.fetchall(
        f"SELECT id FROM people WHERE submitted_by = ? AND {sql_not_approved(db)}",
        (user_id,),
    )
    if not rows:
        rows = db.fetchall("SELECT id FROM people WHERE submitted_by = ?", (user_id,))
    for row in rows:
        pid = int(row["id"])
        db.execute("UPDATE people SET approved = ? WHERE id = ?", (bool_db(db, True), pid))
        place_person_in_tree(db, pid)
    log_audit(db, actor_user_id, "user", user_id, "approve_member", {"people": len(rows)})


def reject_member_registration(db, user_id: int, *, actor_user_id=None):
    """Tolak pendaftaran member."""
    user_id = int(user_id)
    db.execute(
        "UPDATE users SET member_status = 'rejected' WHERE id = ? AND role = 'member'",
        (user_id,),
    )
    log_audit(db, actor_user_id, "user", user_id, "reject_member", {})


def fetch_pending_registrations(db):
    """Member pending beserta data tarombo yang diajukan."""
    return db.fetchall(
        f"""
        SELECT
          u.id AS user_id,
          u.full_name AS account_name,
          u.email,
          u.member_status,
          u.created_at AS registered_at,
          p.id AS person_id,
          p.full_name,
          p.gender,
          p.marga,
          p.tarombo_status,
          p.father_name,
          p.mother_name,
          p.reference_female_line_name,
          p.spouse_name,
          p.spouse_marga,
          (SELECT COUNT(*) FROM submitted_children sc WHERE sc.parent_person_id = p.id) AS child_count
        FROM users u
        LEFT JOIN people p ON p.submitted_by = u.id
        WHERE u.role = 'member' AND u.member_status = 'pending'
        ORDER BY u.created_at DESC
        """
    )


def insert_relationship_ignore(db, source_id, target_id, relation_type):
    if db.backend == "postgres":
        db.execute(
            """
            INSERT INTO relationships (source_person_id, target_person_id, relation_type)
            VALUES (?, ?, ?)
            ON CONFLICT (source_person_id, target_person_id, relation_type) DO NOTHING
            """,
            (source_id, target_id, relation_type),
        )
    else:
        db.execute(
            """
            INSERT OR IGNORE INTO relationships (source_person_id, target_person_id, relation_type)
            VALUES (?, ?, ?)
            """,
            (source_id, target_id, relation_type),
        )


# Akun bawaan tetap: admin, pengurus, anggota, developer — password awal 12345678.
DEFAULT_ACCOUNTS = (
    ("admin@ptsbi.org", "Admin PTSBI", "admin", "12345678", "active"),
    ("pengurus@ptsbi.org", "Pengurus PTSBI", "pengurus", "12345678", "active"),
    ("anggota@ptsbi.org", "Anggota PTSBI", "member", "12345678", "active"),
    ("developer@ptsbi.org", "Developer PTSBI", "developer", "12345678", "active"),
)
DEFAULT_ACCOUNTS_FLAG = "default_accounts_v2"


def _get_app_setting(db, key):
    try:
        row = db.fetchone("SELECT value FROM app_settings WHERE key = ?", (key,))
    except Exception:
        return None
    return row.get("value") if row else None


def _set_app_setting(db, key, value):
    if db.backend == "postgres":
        db.execute(
            """
            INSERT INTO app_settings (key, value, updated_at)
            VALUES (?, ?, NOW())
            ON CONFLICT (key) DO UPDATE SET value = EXCLUDED.value, updated_at = NOW()
            """,
            (key, value),
        )
    else:
        db.execute(
            """
            INSERT OR REPLACE INTO app_settings (key, value, updated_at)
            VALUES (?, ?, CURRENT_TIMESTAMP)
            """,
            (key, value),
        )


def _create_account_if_missing(db, email, name, role, pwd, status):
    if db.fetchone("SELECT id FROM users WHERE email = ?", (email,)):
        return False
    insert_returning_id(
        db,
        """
        INSERT INTO users (full_name, email, password_hash, role, member_status)
        VALUES (?, ?, ?, ?, ?)
        """,
        (name, email, generate_password_hash(pwd), role, status),
    )
    return True


def seed_default_accounts(db):
    """
    Pastikan 4 akun bawaan ada (admin/pengurus/anggota/developer).
    Sekali jalan (ditandai flag) password keempatnya diset ke 12345678 + peran/status
    diselaraskan — tanpa menghapus data lain. Setelah flag aktif, perubahan password
    oleh pengguna tidak akan ditimpa lagi.
    """
    apply_password_reset = _get_app_setting(db, DEFAULT_ACCOUNTS_FLAG) != "1"
    for email, name, role, pwd, status in DEFAULT_ACCOUNTS:
        created = _create_account_if_missing(db, email, name, role, pwd, status)
        if not created and apply_password_reset:
            db.execute(
                """
                UPDATE users SET password_hash = ?, role = ?, member_status = ?
                WHERE email = ?
                """,
                (generate_password_hash(pwd), role, status, email),
            )
    if apply_password_reset:
        _set_app_setting(db, DEFAULT_ACCOUNTS_FLAG, "1")
    db.commit()


def seed_admin_only(db):
    """Akun bawaan tetap (setelah reset database): keempat akun, password 12345678."""
    for email, name, role, pwd, status in DEFAULT_ACCOUNTS:
        _create_account_if_missing(db, email, name, role, pwd, status)
    _set_app_setting(db, DEFAULT_ACCOUNTS_FLAG, "1")
    db.commit()


def seed_data():
    db = get_db()
    seed_default_accounts(db)
    seed_demo_tree(db)


def finalize_new_tarombo_person(db, person_id, *, promote_children=True):
    """Setelah orang approved: tempatkan di silsilah lalu sinkron tampilan."""
    place_person_in_tree(db, person_id, promote_children=promote_children)


def bootstrap_after_db_restore():
    """Buka ulang schema & seed setelah file database diganti."""
    global _db_bootstrapped
    close_database_for_restore()
    bootstrap_database(force=True)
    db = get_db()
    migrate_schema(db)
    backfill_normalized_names(db)


def parse_registration_form(form):
    birth_year_raw = form.get("birth_year", "").strip()
    return {
        "account_full_name": form.get("full_name", "").strip(),
        "email": form.get("email", "").strip().lower(),
        "password": form.get("password", "").strip(),
        "person_name": form.get("person_name", "").strip(),
        "gender": form.get("gender", "").strip(),
        "marga": form.get("marga", "").strip(),
        "birth_year": int(birth_year_raw) if birth_year_raw.isdigit() else None,
        "father_name": form.get("father_name", "").strip(),
        "mother_name": form.get("mother_name", "").strip(),
        "tarombo_status": form.get("tarombo_status", "").strip(),
        "is_married": form.get("is_married", "no").strip(),
        "spouse_name": form.get("spouse_name", "").strip(),
        "spouse_marga": form.get("spouse_marga", "").strip(),
        "reference_female_line_name": form.get("reference_female_line_name", "").strip(),
        "child_names": form.getlist("child_name"),
        "child_genders": form.getlist("child_gender"),
        "child_years": form.getlist("child_birth_year"),
        "child_orders": form.getlist("child_order"),
    }


def register_member_from_request(db, form, *, member_status, approve_immediately, submitted_by_admin_id=None):
    data = parse_registration_form(form)
    if not data["password"] or len(data["password"]) < 6:
        raise ValueError("Password minimal 6 karakter.")
    config_error = validate_submission_from_config(db, form)
    if config_error:
        raise ValueError(config_error)
    error = validate_tarombo_form(
        gender=data["gender"],
        marga=data["marga"],
        tarombo_status=data["tarombo_status"],
        father_name=data["father_name"],
        mother_name=data["mother_name"],
        reference_female_line_name=data["reference_female_line_name"],
        is_married=data["is_married"],
        spouse_name=data["spouse_name"],
        spouse_marga=data["spouse_marga"],
    )
    if error:
        raise ValueError(error)

    result = create_member_with_tarombo(
        db,
        bool_db=bool_db,
        account_full_name=data["account_full_name"],
        email=data["email"],
        password=data["password"],
        person_name=data["person_name"],
        gender=data["gender"],
        marga=data["marga"],
        birth_year=data["birth_year"],
        father_name=data["father_name"],
        mother_name=data["mother_name"],
        tarombo_status=data["tarombo_status"],
        is_married=data["is_married"],
        spouse_name=data["spouse_name"],
        spouse_marga=data["spouse_marga"],
        reference_female_line_name=data["reference_female_line_name"] or None,
        child_names=data["child_names"],
        child_genders=data["child_genders"],
        child_years=data["child_years"],
        child_orders=data.get("child_orders"),
        member_status=member_status,
        approve_immediately=approve_immediately,
        submitted_by_admin_id=submitted_by_admin_id,
        person_identity_fields=person_identity_fields,
        normalize_spouse_fields=normalize_spouse_fields,
        save_submitted_children=save_submitted_children,
        find_existing_person=find_existing_person,
    )
    if approve_immediately:
        finalize_new_tarombo_person(db, result["person_id"])
    return result


def seed_demo_tree(db):
    """Contoh silsilah agar pohon tidak kosong di instalasi baru."""
    row = db.fetchone(f"SELECT COUNT(*) AS c FROM people WHERE {sql_approved(db)}")
    if row and int(row["c"]) > 0:
        return

    admin = db.fetchone("SELECT id FROM users WHERE email = ?", ("admin@ptsbi.org",))
    submitter = admin["id"] if admin else None

    def add_person(name, gender, marga, year, father, mother, status):
        ident = person_identity_fields(name, marga, father, year)
        return insert_returning_id(
            db,
            """
            INSERT INTO people (
              full_name, name_normalized, gender, marga, birth_year,
              father_name, father_name_normalized, mother_name, tarombo_status,
              submitted_by, approved
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                name,
                ident["name_normalized"],
                gender,
                marga,
                year,
                father,
                ident["father_name_normalized"],
                mother,
                status,
                submitter,
                bool_db(db, True),
            ),
        )

    lumban_id = add_person(
        "Lumban Tobing Samosir",
        "male",
        "Samosir",
        1940,
        "Togar Samosir",
        "Siti Samosir",
        "anak",
    )
    hotma_id = add_person(
        "Hotmauli Samosir",
        "male",
        "Samosir",
        1965,
        "Lumban Tobing Samosir",
        "Boru Samosir",
        "anak",
    )
    joni_id = add_person(
        "Joni Samosir",
        "male",
        "Samosir",
        1990,
        "Hotmauli Samosir",
        "Ibu Joni",
        "anak",
    )
    db.commit()
    insert_relationship_ignore(db, lumban_id, hotma_id, "parent")
    insert_relationship_ignore(db, hotma_id, joni_id, "parent")
    db.execute("UPDATE people SET sundut = ?, sundut_locked = ? WHERE id = ?", (1, 1, lumban_id))
    db.execute("UPDATE people SET sundut = ?, sundut_locked = ? WHERE id = ?", (2, 1, hotma_id))
    db.execute("UPDATE people SET sundut = ?, sundut_locked = ? WHERE id = ?", (3, 1, joni_id))
    db.commit()


def get_person_by_name(db, full_name, marga="Samosir", father_name=None, birth_year=None, *, gender=None):
    from services.matching import find_person_loose

    row, _match = find_existing_person(
        db, full_name, marga, father_name=father_name, birth_year=birth_year, approved_only=True
    )
    if row:
        return row
    return find_person_loose(db, full_name, marga, gender=gender, approved_only=True)


def infer_marga_from_name(full_name):
    if not full_name:
        return None
    parts = [p for p in full_name.strip().split() if p]
    return parts[-1] if parts else None


def get_first_token(full_name):
    if not full_name:
        return ""
    return full_name.strip().split()[0]


def format_batak_married_name(person, spouse=None):
    """Alias kompatibilitas — gunakan format_person_display_name."""
    return format_person_display_name(person, spouse)


def get_or_create_parent_person(db, parent_name, gender, marga_hint, child_father_name=None):
    from services.matching import find_person_loose

    marga = infer_marga_from_name(parent_name) or marga_hint or "Samosir"
    existing, _ = find_existing_person(
        db,
        parent_name,
        marga,
        father_name=child_father_name if gender == "male" else None,
        approved_only=True,
    )
    if not existing:
        existing = find_person_loose(db, parent_name, marga, gender=gender, approved_only=True)
    if existing:
        return existing["id"]

    ident = person_identity_fields(parent_name, marga, child_father_name, None)
    return insert_returning_id(
        db,
        """
        INSERT INTO people (
          full_name, name_normalized, gender, marga, birth_year, father_name, father_name_normalized,
          mother_name, tarombo_status, reference_female_line_name, spouse_name, spouse_marga,
          submitted_by, approved
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            parent_name,
            ident["name_normalized"],
            gender,
            marga,
            None,
            "Tidak Diketahui",
            ident["father_name_normalized"],
            "Tidak Diketahui",
            "anak",
            None,
            None,
            None,
            None,
            bool_db(db, True),
        ),
    )


def create_parent_edge(db, child_id, parent_name, parent_gender, child_marga, child_father_name=None):
    if not parent_name or parent_name.strip().lower() in {"tidak diketahui", "unknown", "-"}:
        return
    parent_id = get_or_create_parent_person(
        db, parent_name, parent_gender, child_marga, child_father_name=child_father_name
    )
    if parent_id:
        insert_relationship_ignore(db, parent_id, child_id, "parent")


def build_relationships_for_person(db, person_id):
    person = db.fetchone(
        f"SELECT id, father_name, mother_name, marga FROM people WHERE id = ? AND {sql_approved(db)}",
        (person_id,),
    )
    if not person:
        return
    # Garis pohon hanya mengikuti ayah (laki-laki); ibu dihubungkan sebagai pasangan, bukan parent→anak.
    create_parent_edge(
        db, person["id"], person["father_name"], "male", person["marga"], child_father_name=person["father_name"]
    )

    father = get_person_by_name(db, person["father_name"], person["marga"])
    mother = get_person_by_name(db, person["mother_name"], "Samosir")
    if father and mother:
        insert_relationship_ignore(db, father["id"], mother["id"], "spouse")
        insert_relationship_ignore(db, mother["id"], father["id"], "spouse")
    link_spouse_from_fields(db, person_id)


def link_spouse_from_fields(db, person_id):
    """Hubungkan pasangan (istri/suami) lewat relasi spouse agar tampil di pohon."""
    person = db.fetchone(f"SELECT * FROM people WHERE id = ? AND {sql_approved(db)}", (person_id,))
    if not person:
        return
    spouse_name = (person.get("spouse_name") or "").strip()
    if not spouse_name or spouse_name.lower() in {"tidak diketahui", "unknown", "-"}:
        return

    spouse_marga = person.get("spouse_marga") or infer_marga_from_name(spouse_name) or "Samosir"
    spouse_gender = "female" if person["gender"] == "male" else "male"
    spouse_row = get_person_by_name(db, spouse_name, spouse_marga, gender=spouse_gender)

    if person["gender"] == "male":
        if spouse_row:
            spouse_id = spouse_row["id"]
            spouse_marga = person.get("spouse_marga") or spouse_row.get("marga") or spouse_marga
            if (person.get("marga") or "").lower() == "samosir":
                wife_marga = (spouse_row.get("marga") or spouse_marga or "").strip().lower()
                wife_status = "boru" if wife_marga == "samosir" else "anak"
                db.execute(
                    """
                    UPDATE people SET tarombo_status = ?, spouse_name = ?, spouse_marga = ?
                    WHERE id = ?
                    """,
                    (wife_status, person["full_name"], person["marga"], spouse_id),
                )
        else:
            if (person.get("marga") or "").lower() == "samosir":
                sm = (spouse_marga or "").strip().lower()
                spouse_status = "boru" if sm == "samosir" else "anak"
            else:
                spouse_status = "bere"
            ident = person_identity_fields(spouse_name, spouse_marga, None, None)
            spouse_id = insert_returning_id(
                db,
                """
                INSERT INTO people (
                  full_name, name_normalized, gender, marga, birth_year, father_name, father_name_normalized,
                  mother_name, tarombo_status, reference_female_line_name, spouse_name, spouse_marga,
                  submitted_by, approved
                ) VALUES (?, ?, 'female', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    spouse_name,
                    ident["name_normalized"],
                    spouse_marga,
                    None,
                    "Tidak Diketahui",
                    ident["father_name_normalized"],
                    "Tidak Diketahui",
                    spouse_status,
                    None,
                    person["full_name"],
                    person["marga"],
                    coerce_submitted_by(db, person.get("submitted_by")),
                    bool_db(db, True),
                ),
            )
    elif person["gender"] == "female" and spouse_row and spouse_row["gender"] == "male":
        spouse_id = spouse_row["id"]
        db.execute(
            """
            UPDATE people SET spouse_name = ?, spouse_marga = ?
            WHERE id = ? AND (spouse_name IS NULL OR TRIM(spouse_name) = '')
            """,
            (person["full_name"], person["marga"], spouse_id),
        )
    else:
        return

    insert_relationship_ignore(db, person_id, spouse_id, "spouse")
    insert_relationship_ignore(db, spouse_id, person_id, "spouse")


def save_submitted_children(db, parent_person_id, child_names, child_genders, child_years, child_orders=None):
    child_orders = child_orders or []
    for idx, raw_name in enumerate(child_names):
        child_name = raw_name.strip()
        if not child_name:
            continue
        child_gender = child_genders[idx] if idx < len(child_genders) else "male"
        if child_gender not in ("male", "female"):
            child_gender = "male"
        year_raw = child_years[idx].strip() if idx < len(child_years) and child_years[idx] else ""
        child_year = int(year_raw) if year_raw.isdigit() else None
        order_raw = child_orders[idx] if idx < len(child_orders) else ""
        child_order = parse_child_order(order_raw)
        db.execute(
            """
            INSERT INTO submitted_children (parent_person_id, child_name, child_gender, child_birth_year, child_order)
            VALUES (?, ?, ?, ?, ?)
            """,
            (parent_person_id, child_name, child_gender, child_year, child_order),
        )


def resolve_parents_for_child(parent):
    if parent["gender"] == "male":
        return parent["full_name"], parent.get("spouse_name") or parent["mother_name"]
    return parent.get("spouse_name") or parent["father_name"], parent["full_name"]


def child_status_and_marga(parent, child_name, child_gender):
    if child_gender == "male":
        return parent["marga"], "anak"
    child_marga = infer_marga_from_name(child_name) or parent["marga"]
    if parent["gender"] == "male" and parent["marga"].lower() == "samosir":
        child_marga = "Samosir"
    status = "boru" if child_marga.lower() == "samosir" else "bere"
    return child_marga, status


def add_child_for_parent(
    db, parent_person_id, child_name, child_gender, child_birth_year, approved=True, child_order=None
):
    from services.matching import find_person_loose

    parent = db.fetchone("SELECT * FROM people WHERE id = ?", (parent_person_id,))
    if not parent or not child_name.strip():
        return None
    father_name, mother_name = resolve_parents_for_child(parent)
    child_marga, child_status = child_status_and_marga(parent, child_name, child_gender)
    existing_child = find_person_loose(
        db, child_name.strip(), child_marga, gender=child_gender, approved_only=approved
    )
    if existing_child:
        if approved:
            create_parent_edge(db, existing_child["id"], father_name, "male", parent["marga"], child_father_name=father_name)
            create_parent_edge(db, existing_child["id"], mother_name, "female", parent["marga"])
        return existing_child["id"]
    year = int(child_birth_year) if child_birth_year and str(child_birth_year).isdigit() else None
    order = parse_child_order(child_order)
    ident = person_identity_fields(child_name.strip(), child_marga, father_name, year)
    child_id = insert_returning_id(
        db,
        """
        INSERT INTO people (
          full_name, name_normalized, gender, marga, birth_year, child_order, father_name, father_name_normalized,
          mother_name, tarombo_status, reference_female_line_name, spouse_name, spouse_marga,
          submitted_by, approved
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            child_name.strip(),
            ident["name_normalized"],
            child_gender,
            child_marga,
            year,
            order,
            father_name,
            ident["father_name_normalized"],
            mother_name,
            child_status,
            None,
            None,
            None,
            coerce_submitted_by(db, parent.get("submitted_by")),
            bool_db(db, approved),
        ),
    )
    if approved:
        create_parent_edge(db, child_id, father_name, "male", parent["marga"], child_father_name=father_name)
        create_parent_edge(db, child_id, mother_name, "female", parent["marga"])
    return child_id


def get_children_of_parent(db, parent_person_id):
    rows = db.fetchall(
        """
        SELECT p.*
        FROM relationships r
        JOIN people p ON p.id = r.target_person_id
        WHERE r.source_person_id = ? AND r.relation_type = 'parent'
        """,
        (parent_person_id,),
    )
    return sort_people_list([dict(r) for r in rows])


def get_submitted_children_of_parent(db, parent_person_id):
    rows = db.fetchall(
        "SELECT * FROM submitted_children WHERE parent_person_id = ?",
        (parent_person_id,),
    )
    return sort_people_list(
        [
            {
                **dict(r),
                "birth_year": r.get("child_birth_year"),
                "child_order": r.get("child_order"),
            }
            for r in rows
        ]
    )


def create_children_for_person(db, parent_person_id):
    children = get_submitted_children_of_parent(db, parent_person_id)
    for child in children:
        add_child_for_parent(
            db,
            parent_person_id,
            child["child_name"],
            child["child_gender"],
            child["child_birth_year"],
            approved=True,
            child_order=child.get("child_order"),
        )
    db.execute("DELETE FROM submitted_children WHERE parent_person_id = ?", (parent_person_id,))


def rebuild_all_parent_relationships(db):
    approved_people = db.fetchall(f"SELECT id FROM people WHERE {sql_approved(db)} ORDER BY id ASC")
    db.execute("DELETE FROM relationships WHERE relation_type = 'parent'")
    for person in approved_people:
        build_relationships_for_person(db, person["id"])
    for person in approved_people:
        link_spouse_from_fields(db, person["id"])
    db.commit()


def fix_spouse_boru_status(db):
    """
    Selaraskan status pasangan dengan aturan marga Samosir:
    istri non-Samosir dari laki-laki Samosir = anak; perempuan marga Samosir = boru.
    """
    pairs = db.fetchall(
        f"""
        SELECT r.source_person_id AS husband_id, r.target_person_id AS wife_id
        FROM relationships r
        JOIN people h ON h.id = r.source_person_id AND {sql_approved(db, 'h')}
        JOIN people w ON w.id = r.target_person_id AND {sql_approved(db, 'w')}
        WHERE r.relation_type = 'spouse' AND h.gender = 'male' AND LOWER(h.marga) = 'samosir'
        """
    )
    for row in pairs:
        wife = db.fetchone(
            f"SELECT marga, tarombo_status FROM people WHERE id = ? AND {sql_approved(db)}",
            (row["wife_id"],),
        )
        if not wife:
            continue
        wife_status = "boru" if (wife.get("marga") or "").strip().lower() == "samosir" else "anak"
        db.execute(
            """
            UPDATE people SET tarombo_status = ?
            WHERE id = ? AND gender = 'female'
            """,
            (wife_status, row["wife_id"]),
        )
    db.execute(
        f"""
        UPDATE people SET tarombo_status = 'boru'
        WHERE gender = 'female' AND LOWER(marga) = 'samosir'
          AND tarombo_status = 'anak' AND {sql_approved(db)}
        """
    )
    db.commit()


def resync_graph(db):
    """Sinkron relasi ayah, panggoaran, sundut (tanpa gabung duplikat)."""
    sync_tarombo(
        db,
        sql_approved,
        insert_relationship_ignore,
        merge=False,
        purge=False,
    )


def calculate_generations(people):
    """Fallback generasi dari rantai nama ayah (hanya bila sundut DB kosong)."""
    by_name = {}
    by_id = {}
    for person in people:
        by_id[person["id"]] = person
        if person.get("full_name"):
            by_name[person["full_name"]] = person
    generation_cache = {}

    def get_gen(person_id):
        if person_id in generation_cache:
            return generation_cache[person_id]

        path = []
        current_id = person_id
        while current_id is not None:
            if current_id in generation_cache:
                base_gen = generation_cache[current_id]
                if base_gen is None:
                    for pid in path:
                        generation_cache[pid] = None
                    return None
                for offset, pid in enumerate(reversed(path)):
                    generation_cache[pid] = base_gen + offset + 1
                return generation_cache[person_id]

            if current_id in path:
                for pid in path:
                    generation_cache[pid] = None
                generation_cache[current_id] = None
                return None

            path.append(current_id)
            person = by_id.get(current_id)
            if not person:
                for pid in path:
                    generation_cache[pid] = None
                return None

            father_name = person.get("father_name")
            father = by_name.get(father_name) if father_name else None
            if not father or father["id"] == current_id:
                for offset, pid in enumerate(path):
                    generation_cache[pid] = offset + 1
                return generation_cache[person_id]

            current_id = father["id"]
            if len(path) > 50:
                for pid in path:
                    generation_cache[pid] = None
                return None

        return generation_cache.get(person_id)

    for person in people:
        get_gen(person["id"])
    return generation_cache


def _redirect_legacy_hasil(*, layar_penuh: bool):
    """URL /hasil* lama → /tarombo (bookmark)."""
    if not session.get("user_id"):
        flash("Silakan masuk untuk melihat Tarombo.", "error")
        return redirect(url_for("login"))
    if not user_can_view_tarombo():
        flash("Tarombo hanya untuk member yang sudah disetujui (aktif).", "error")
        return redirect(url_for("dashboard"))
    if layar_penuh:
        return redirect(url_for("tarombo_layar_penuh"))
    return redirect(url_for("tarombo"))


@app.route("/hasil")
def hasil_legacy_redirect():
    return _redirect_legacy_hasil(layar_penuh=False)


@app.route("/hasil/layar-penuh")
def hasil_layar_legacy_redirect():
    return _redirect_legacy_hasil(layar_penuh=True)


@app.route("/api/ping")
def api_ping():
    return jsonify({"ok": True})


@app.route("/api/fix-db")
@developer_required
def api_fix_db():
    """Perbaiki DB lokal tanpa restart (dev)."""
    global _db_bootstrapped
    try:
        reset_request_db()
        reset_sqlite_database_file()
        _db_bootstrapped = False
        bootstrap_database(force=True)
        db = get_db()
        payload = build_heritage_tree_payload(db)
        return jsonify(
            {
                "ok": True,
                "roots": len(payload["roots"]),
                "people": payload["stats"]["total_people"],
            }
        )
    except Exception as exc:
        app.logger.exception("api_fix_db")
        return jsonify({"ok": False, "error": str(exc)}), 500


@app.route("/")
def home():
    return render_template("index.html")


@app.route("/register", methods=["GET", "POST"])
def register():
    if session.get("user_id"):
        return redirect(url_for("dashboard"))
    db = get_db()
    if request.method == "POST":
        try:
            register_member_from_request(
                db,
                request.form,
                member_status="pending",
                approve_immediately=False,
            )
        except ValueError as exc:
            flash(str(exc), "error")
            return render_template("register.html", **register_form_context(db))
        flash("Pendaftaran berhasil. Menunggu verifikasi admin.", "success")
        return redirect(url_for("login"))
    return render_template("register.html", **register_form_context(db))


def normalize_spouse_fields(gender, marga, is_married, spouse_name, spouse_marga):
    spouse_name = (spouse_name or "").strip()
    spouse_marga = (spouse_marga or "").strip()
    if is_married != "yes":
        return None, None
    if gender == "male" and marga.lower() == "samosir":
        return spouse_name or None, spouse_marga or None
    if gender == "female":
        return spouse_name or None, spouse_marga or None
    return spouse_name or None, spouse_marga or None


def validate_tarombo_form(
    gender,
    marga,
    tarombo_status,
    father_name,
    mother_name,
    reference_female_line_name,
    is_married,
    spouse_name,
    spouse_marga,
):
    if not father_name or not mother_name:
        return "Nama ayah dan ibu wajib diisi agar generasi tarombo dapat dihitung."
    if gender == "male" and marga.lower() == "samosir" and is_married == "yes":
        if not spouse_name:
            return "Nama pasangan wajib diisi untuk pria marga Samosir yang sudah menikah."
        if not spouse_marga:
            return "Marga pasangan wajib diisi untuk pria marga Samosir yang sudah menikah."
    if tarombo_status == "anak" and gender != "male":
        return "Status Anak harus laki-laki."
    if tarombo_status == "boru":
        if gender != "female":
            return "Status Boru harus perempuan."
        if marga.lower() != "samosir":
            return "Status Boru wajib bermarga Samosir."
        if is_married == "yes":
            if not spouse_name or not spouse_marga:
                return "Jika Boru menikah, nama dan marga suami wajib diisi."
            if spouse_marga.lower() == "samosir":
                return "Suami Boru harus non-Samosir sesuai aturan yang ditetapkan."
    if tarombo_status in ("bere", "ibebere"):
        if not reference_female_line_name:
            return "Bere/Ibebere wajib mengisi nama saudara perempuan marga Samosir sebagai acuan garis."
        if is_married == "yes" and not spouse_name:
            return "Jika Bere/Ibebere menikah, nama pasangan wajib diisi."
    return None


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "").strip()
        user = get_db().fetchone("SELECT * FROM users WHERE email = ?", (email,))
        if not user or not check_password_hash(user["password_hash"], password):
            flash("Email atau password salah.", "error")
            return render_template("login.html")
        session["user_id"] = user["id"]
        session["role"] = user["role"]
        session["member_status"] = user["member_status"]
        session["full_name"] = user["full_name"]
        if is_staff_role(user["role"]):
            return redirect(url_for(staff_redirect_endpoint()))
        return redirect(url_for("dashboard"))
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("home"))


@app.route("/dashboard")
@login_required
def dashboard():
    return render_template("member_dashboard.html")


@app.route("/account/password", methods=["GET", "POST"])
@login_required
def change_password():
    back_url = url_for("admin") if is_staff_role() else url_for("dashboard")
    if request.method == "POST":
        current = request.form.get("current_password", "").strip()
        new_pwd = request.form.get("new_password", "").strip()
        confirm = request.form.get("confirm_password", "").strip()
        if not current or not new_pwd or not confirm:
            flash("Semua kolom password wajib diisi.", "error")
            return render_template("change_password.html", back_url=back_url)
        if len(new_pwd) < 6:
            flash("Password baru minimal 6 karakter.", "error")
            return render_template("change_password.html", back_url=back_url)
        if new_pwd != confirm:
            flash("Konfirmasi password baru tidak sama.", "error")
            return render_template("change_password.html", back_url=back_url)
        db = get_db()
        user = db.fetchone("SELECT id, password_hash FROM users WHERE id = ?", (session["user_id"],))
        if not user or not check_password_hash(user["password_hash"], current):
            flash("Password saat ini salah.", "error")
            return render_template("change_password.html", back_url=back_url)
        db.execute(
            "UPDATE users SET password_hash = ? WHERE id = ?",
            (generate_password_hash(new_pwd), session["user_id"]),
        )
        db.commit()
        flash("Password berhasil diubah.", "success")
        return redirect(back_url)
    return render_template("change_password.html", back_url=back_url)


def _member_own_person(db, user_id):
    """Record tarombo milik anggota sendiri (submitted_by = user_id)."""
    return db.fetchone(
        "SELECT * FROM people WHERE submitted_by = ? ORDER BY id ASC LIMIT 1",
        (user_id,),
    )


@app.route("/profile", methods=["GET", "POST"])
@login_required
def profile():
    """Anggota melengkapi / mengubah data akun & data tarombo miliknya sendiri."""
    db = get_db()
    user_id = session["user_id"]
    person = _member_own_person(db, user_id)

    if request.method == "POST":
        action = request.form.get("action", "save_profile")

        if action == "save_account":
            full_name = request.form.get("account_full_name", "").strip()
            if not full_name:
                flash("Nama akun wajib diisi.", "error")
                return redirect(url_for("profile"))
            db.execute("UPDATE users SET full_name = ? WHERE id = ?", (full_name, user_id))
            db.commit()
            session["full_name"] = full_name
            flash("Nama akun diperbarui.", "success")
            return redirect(url_for("profile"))

        # action == save_profile (data tarombo milik sendiri)
        person_name = request.form.get("full_name", "").strip()
        gender = request.form.get("gender", "").strip()
        marga = request.form.get("marga", "").strip()
        tarombo_status = request.form.get("tarombo_status", "").strip()
        father_name = request.form.get("father_name", "").strip()
        mother_name = request.form.get("mother_name", "").strip()
        reference_female_line_name = request.form.get("reference_female_line_name", "").strip()
        spouse_name_raw = request.form.get("spouse_name", "").strip()
        spouse_marga_raw = request.form.get("spouse_marga", "").strip()
        is_married = "yes" if spouse_name_raw else "no"

        if not person_name or not gender or not marga or not tarombo_status:
            flash("Nama, jenis kelamin, marga, dan status tarombo wajib diisi.", "error")
            return redirect(url_for("profile"))

        error = validate_tarombo_form(
            gender=gender,
            marga=marga,
            tarombo_status=tarombo_status,
            father_name=father_name,
            mother_name=mother_name,
            reference_female_line_name=reference_female_line_name,
            is_married=is_married,
            spouse_name=spouse_name_raw,
            spouse_marga=spouse_marga_raw,
        )
        if error:
            flash(error, "error")
            return redirect(url_for("profile"))

        year_raw = request.form.get("birth_year", "").strip()
        year = int(year_raw) if year_raw.isdigit() else None
        spouse_name, spouse_marga = normalize_spouse_fields(
            gender, marga, is_married, spouse_name_raw, spouse_marga_raw
        )
        ident = person_identity_fields(person_name, marga, father_name, year, mother_name)

        if person:
            db.execute(
                """
                UPDATE people SET
                  full_name = ?, name_normalized = ?, gender = ?, marga = ?, birth_year = ?,
                  father_name = ?, father_name_normalized = ?, mother_name = ?, mother_name_normalized = ?,
                  tarombo_status = ?, reference_female_line_name = ?, spouse_name = ?, spouse_marga = ?
                WHERE id = ?
                """,
                (
                    person_name,
                    ident["name_normalized"],
                    gender,
                    marga,
                    year,
                    father_name,
                    ident["father_name_normalized"],
                    mother_name,
                    ident["mother_name_normalized"],
                    tarombo_status,
                    reference_female_line_name or None,
                    spouse_name,
                    spouse_marga,
                    person["id"],
                ),
            )
            person_id = person["id"]
            already_approved = bool(person.get("approved"))
        else:
            person_id = insert_returning_id(
                db,
                """
                INSERT INTO people (
                  full_name, name_normalized, gender, marga, birth_year,
                  father_name, father_name_normalized, mother_name, mother_name_normalized,
                  tarombo_status, reference_female_line_name, spouse_name, spouse_marga,
                  submitted_by, approved
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    person_name,
                    ident["name_normalized"],
                    gender,
                    marga,
                    year,
                    father_name,
                    ident["father_name_normalized"],
                    mother_name,
                    ident["mother_name_normalized"],
                    tarombo_status,
                    reference_female_line_name or None,
                    spouse_name,
                    spouse_marga,
                    user_id,
                    bool_db(db, False),
                ),
            )
            already_approved = False

        db.commit()
        if already_approved:
            try:
                place_person_in_tree(db, person_id, promote_children=False)
                db.commit()
            except Exception:
                app.logger.exception("place_person_in_tree (profile) gagal")
        log_audit(db, user_id, "person", person_id, "profile_save", {"approved": already_approved})
        account_row = db.fetchone("SELECT email FROM users WHERE id = ?", (user_id,))
        if account_row and account_row.get("email"):
            upsert_member_profile_record(
                db,
                email=account_row["email"],
                profile={
                    "source": "tarombo",
                    "person_id": person_id,
                    "full_name": person_name,
                    "marga": marga,
                    "tarombo_status": tarombo_status,
                    "gender": gender,
                },
                wp_user_id=None,
            )
        db.commit()
        if already_approved:
            flash("Profil tersimpan dan diperbarui di pohon tarombo.", "success")
        else:
            flash("Profil tersimpan. Menunggu persetujuan admin untuk tampil di pohon.", "success")
        return redirect(url_for("profile"))

    account = db.fetchone(
        "SELECT full_name, email, role, member_status FROM users WHERE id = ?",
        (user_id,),
    )
    return render_template(
        "profile.html",
        account=account,
        person=person,
        partuturan_options=PARTUTURAN_STATUSES,
    )


@app.route("/tarombo")
@login_required
def tarombo():
    if not user_can_view_tarombo():
        flash("Tarombo hanya dapat diakses setelah admin menyetujui pendaftaran Anda.", "error")
        return redirect(url_for("dashboard"))
    db = get_db()
    payload = build_heritage_tree_payload(db)
    people = db.fetchall(
        f"""
        SELECT id, full_name, marga, gender, tarombo_status, father_name, mother_name, spouse_name, spouse_marga, submitted_by
        FROM people
        WHERE {sql_approved(db)} AND full_name != 'Tidak Diketahui'
        ORDER BY created_at DESC
        """
    )
    return render_template(
        "tarombo.html",
        people=people,
        tree_json=json.dumps(payload["roots"]),
        stats_json=json.dumps(payload["stats"]),
        index_json=json.dumps(payload["index"]),
    )


@app.route("/tarombo/layar-penuh")
@login_required
def tarombo_layar_penuh():
    if not user_can_view_tarombo():
        flash("Tarombo hanya dapat diakses setelah admin menyetujui pendaftaran Anda.", "error")
        return redirect(url_for("dashboard"))
    db = get_db()
    payload = build_heritage_tree_payload(db)
    return render_template(
        "tarombo_layar_penuh.html",
        tree_json=json.dumps(payload["roots"], ensure_ascii=False, default=str),
        stats_json=json.dumps(payload["stats"], ensure_ascii=False, default=str),
        index_json=json.dumps(payload["index"], ensure_ascii=False, default=str),
    )


def build_heritage_tree_payload(db, *, sync=False):
    try:
        fix_spouse_boru_status(db)
    except Exception:
        app.logger.exception("fix_spouse_boru_status gagal")
    roots = build_collapsible_tree(db, sync=sync)
    tree_ids = collect_tree_person_ids(roots)
    if tree_ids:
        placeholders = ",".join("?" * len(tree_ids))
        people = db.fetchall(
            f"""
            SELECT id, full_name, marga, gender, birth_year, father_name, mother_name, spouse_name, spouse_marga, panggoaran
            FROM people
            WHERE {sql_approved(db)} AND full_name != 'Tidak Diketahui' AND id IN ({placeholders})
            ORDER BY full_name ASC
            """,
            tuple(tree_ids),
        )
    else:
        people = []
    index = [{"id": p["id"], "name": p["full_name"], "marga": p["marga"], "birth_year": p["birth_year"]} for p in people]
    max_gen = 0

    def walk_depth(node, depth=1):
        nonlocal max_gen
        max_gen = max(max_gen, depth)
        for child in node.get("children", []):
            walk_depth(child, depth + 1)

    for root in roots:
        walk_depth(root, 1)

    max_sundut = db.fetchone(
        f"SELECT MAX(sundut) AS m FROM people WHERE {sql_approved(db)} AND sundut IS NOT NULL"
    )
    max_from_db = int(max_sundut["m"]) if max_sundut and max_sundut.get("m") is not None else 0
    return {
        "roots": roots,
        "stats": {
            "total_people": len(people),
            "root_families": len(roots),
            "max_generation": max(max_gen, max_from_db),
        },
        "index": index,
    }


def build_collapsible_tree(db, *, sync=False):
    if sync:
        resync_graph(db)
    people = db.fetchall(
        f"""
        SELECT id, full_name, marga, gender, birth_year, child_order, tarombo_status, father_name, mother_name,
               spouse_name, spouse_marga, submitted_by,
               panggoaran, panggoaran_type, sundut, sundut_locked
        FROM people
        WHERE {sql_approved(db)} AND full_name != 'Tidak Diketahui'
        ORDER BY id ASC
        """
    )
    if not people:
        return []

    allowed_ids = {p["id"] for p in people}
    all_rels = db.fetchall("SELECT source_person_id, target_person_id, relation_type FROM relationships")
    relationships = [
        r
        for r in all_rels
        if r["source_person_id"] in allowed_ids and r["target_person_id"] in allowed_ids
    ]

    people = enrich_people_panggoaran_for_tree(people, relationships)

    return build_unique_tree(
        people,
        relationships,
        format_batak_married_name=format_batak_married_name,
        format_panggoaran_display=format_panggoaran_display,
        get_opung_hover_suffix=get_opung_hover_suffix,
        effective_sundut=lambda person, fb: effective_sundut(person, fb),
        fallback_generations=calculate_generations(people),
        sundut_labels=sundut_title_map(db),
    )


@app.route("/api/tarombo-tree")
@login_required
def tarombo_tree_api():
    if not user_can_view_tarombo():
        return jsonify({"error": "forbidden"}), 403
    return jsonify(build_heritage_tree_payload(get_db(), sync=False))


@app.route("/admin/pohon")
@login_required
@staff_required
def admin_pohon():
    """Editor pohon: drag-drop penempatan + panel edit + simpan."""
    db = get_db()
    payload = build_heritage_tree_payload(db)
    return render_template(
        "admin_pohon.html",
        tree_json=json.dumps(payload["roots"], ensure_ascii=False, default=str),
        stats_json=json.dumps(payload["stats"], ensure_ascii=False, default=str),
        index_json=json.dumps(payload["index"], ensure_ascii=False, default=str),
        partuturan_options=PARTUTURAN_STATUSES,
    )


@app.route("/api/admin/person/<int:person_id>")
@login_required
@staff_required
def api_admin_person(person_id):
    person = get_db().fetchone("SELECT * FROM people WHERE id = ?", (person_id,))
    if not person:
        return jsonify({"error": "not_found"}), 404
    row = dict(person)
    for key in ("approved", "sundut_locked"):
        if key in row:
            row[key] = bool(row[key]) if row[key] is not None else False
    row["partuturan_label"] = partuturan_label(row.get("tarombo_status"))
    return jsonify(row)


@app.route("/api/admin/tree-changes", methods=["POST"])
@login_required
@staff_required
def api_admin_tree_changes():
    data = request.get_json(silent=True) or {}
    changes = data.get("changes") or []
    if not isinstance(changes, list) or not changes:
        return jsonify({"error": "Tidak ada perubahan untuk disimpan."}), 400
    db = get_db()
    try:
        applied = apply_tree_changes(db, changes)
        db.commit()
        for person_id in {a.get("person_id") for a in applied if a.get("person_id")}:
            if person_id:
                build_relationships_for_person(db, int(person_id))
        for item in applied:
            if item.get("op") == "swap_parents":
                build_relationships_for_person(db, int(item["person_id_a"]))
                build_relationships_for_person(db, int(item["person_id_b"]))
            elif item.get("op") == "replace_with":
                build_relationships_for_person(db, int(item["source_id"]))
                build_relationships_for_person(db, int(item["target_id"]))
        db.commit()
        resync_graph(db)
        log_audit(db, session.get("user_id"), "tree_editor", None, "apply", {"count": len(applied)})
        db.commit()
        payload = build_heritage_tree_payload(db, sync=False)
        return jsonify({"ok": True, "applied": applied, "tree": payload})
    except Exception as exc:
        db.rollback()
        app.logger.exception("api_admin_tree_changes")
        return jsonify({"error": str(exc)}), 400


@app.route("/admin/person/<int:person_id>/edit", methods=["GET", "POST"])
@login_required
@staff_required
def admin_edit_person(person_id):
    db = get_db()
    person = db.fetchone("SELECT * FROM people WHERE id = ?", (person_id,))
    if not person:
        flash("Data tidak ditemukan.", "error")
        return redirect(url_for("admin"))

    if request.method == "POST":
        action = request.form.get("action", "save_person")
        if action == "add_child":
            add_child_for_parent(
                db,
                person_id,
                request.form.get("child_name", ""),
                request.form.get("child_gender", "male"),
                request.form.get("child_birth_year", ""),
                approved=True,
                child_order=request.form.get("child_order"),
            )
            db.commit()
            resync_graph(db)
            db.commit()
            flash("Anak berhasil ditambahkan ke tarombo.", "success")
        elif action == "resync":
            resync_graph(db)
            log_audit(db, session.get("user_id"), "person", person_id, "resync", {})
            db.commit()
            flash("Graf silsilah, panggoaran, dan sundut berhasil disinkronkan ulang.", "success")
        elif action == "add_submitted_child":
            save_submitted_children(
                db,
                person_id,
                [request.form.get("child_name", "")],
                [request.form.get("child_gender", "male")],
                [request.form.get("child_birth_year", "")],
                [request.form.get("child_order", "")],
            )
            db.commit()
            flash("Anak (draft) berhasil ditambahkan.", "success")
        elif action == "promote_submitted_child":
            submitted_id = request.form.get("submitted_child_id")
            row = db.fetchone(
                "SELECT * FROM submitted_children WHERE id = ? AND parent_person_id = ?",
                (submitted_id, person_id),
            )
            if row:
                add_child_for_parent(
                    db,
                    person_id,
                    row["child_name"],
                    row["child_gender"],
                    row["child_birth_year"],
                    approved=True,
                    child_order=row.get("child_order"),
                )
                db.execute("DELETE FROM submitted_children WHERE id = ?", (submitted_id,))
                db.commit()
                resync_graph(db)
                flash("Anak draft berhasil dimasukkan ke tarombo.", "success")
        elif action == "delete_submitted_child":
            db.execute("DELETE FROM submitted_children WHERE id = ? AND parent_person_id = ?", (request.form.get("submitted_child_id"), person_id))
            db.commit()
            flash("Anak draft dihapus.", "success")
        elif action == "save_sibling_orders":
            sibling_ids = request.form.getlist("sibling_id")
            sibling_orders = request.form.getlist("sibling_order")
            order_map = {}
            for sid, ord_raw in zip(sibling_ids, sibling_orders):
                if sid and sid.isdigit():
                    order_map[int(sid)] = parse_child_order(ord_raw)
            apply_child_orders(db, {k: v for k, v in order_map.items() if v is not None})
            db.commit()
            resync_graph(db)
            flash("Urutan saudara (satu ayah) berhasil disimpan.", "success")
        elif action == "auto_sibling_orders":
            siblings = get_siblings_for_person(db, person_id, sql_approved, get_person_by_name)
            updates = assign_orders_from_birth_year(siblings)
            if updates:
                apply_child_orders(db, updates)
                db.commit()
                resync_graph(db)
                flash(f"Urutan otomatis dari tahun lahir: {len(updates)} orang diperbarui.", "success")
            else:
                flash("Semua saudara sudah punya urutan manual.", "success")
        elif action == "update_child":
            child_id = request.form.get("child_id")
            parent = db.fetchone("SELECT * FROM people WHERE id = ?", (person_id,))
            father_name, mother_name = resolve_parents_for_child(parent)
            child_name = request.form.get("child_name", "").strip()
            child_gender = request.form.get("child_gender", "male")
            child_marga, child_status = child_status_and_marga(parent, child_name, child_gender)
            ident = person_identity_fields(child_name, child_marga, father_name, None)
            db.execute(
                """
                UPDATE people SET
                  full_name = ?, name_normalized = ?, gender = ?, marga = ?, birth_year = ?,
                  child_order = ?, father_name = ?, father_name_normalized = ?, mother_name = ?, tarombo_status = ?
                WHERE id = ?
                """,
                (
                    child_name,
                    ident["name_normalized"],
                    child_gender,
                    child_marga,
                    int(request.form["child_birth_year"]) if request.form.get("child_birth_year", "").strip().isdigit() else None,
                    parse_child_order(request.form.get("child_order")),
                    father_name,
                    ident["father_name_normalized"],
                    mother_name,
                    child_status,
                    child_id,
                ),
            )
            db.commit()
            resync_graph(db)
            flash("Data anak berhasil diperbarui.", "success")
        elif action == "delete_child":
            child_id = request.form.get("child_id")
            db.execute("DELETE FROM relationships WHERE source_person_id = ? OR target_person_id = ?", (child_id, child_id))
            db.execute("DELETE FROM people WHERE id = ?", (child_id,))
            db.commit()
            resync_graph(db)
            flash("Data anak berhasil dihapus.", "success")
        else:
            gender = request.form.get("gender", "").strip()
            marga = request.form.get("marga", "").strip()
            is_married = "yes" if request.form.get("spouse_name", "").strip() else "no"
            spouse_name, spouse_marga = normalize_spouse_fields(
                gender,
                marga,
                is_married,
                request.form.get("spouse_name", ""),
                request.form.get("spouse_marga", ""),
            )
            year = int(request.form["birth_year"]) if request.form.get("birth_year", "").strip().isdigit() else None
            ident = person_identity_fields(request.form.get("full_name", "").strip(), marga, request.form.get("father_name", "").strip(), year)
            approved = request.form.get("approved") == "1"
            sundut_raw = request.form.get("sundut", "").strip()
            sundut_val = int(sundut_raw) if sundut_raw.isdigit() else None
            sundut_locked = request.form.get("sundut_locked") == "1"
            db.execute(
                """
                UPDATE people SET
                  full_name = ?, name_normalized = ?, gender = ?, marga = ?, birth_year = ?,
                  child_order = ?, father_name = ?, father_name_normalized = ?, mother_name = ?, tarombo_status = ?,
                  reference_female_line_name = ?, spouse_name = ?, spouse_marga = ?,
                  approved = ?, sundut = ?, sundut_locked = ?
                WHERE id = ?
                """,
                (
                    request.form.get("full_name", "").strip(),
                    ident["name_normalized"],
                    gender,
                    marga,
                    year,
                    parse_child_order(request.form.get("child_order")),
                    request.form.get("father_name", "").strip(),
                    ident["father_name_normalized"],
                    request.form.get("mother_name", "").strip(),
                    request.form.get("tarombo_status", "").strip(),
                    request.form.get("reference_female_line_name", "").strip() or None,
                    spouse_name,
                    spouse_marga,
                    bool_db(db, approved),
                    sundut_val,
                    bool_db(db, sundut_locked),
                    person_id,
                ),
            )
            db.commit()
            if approved:
                place_person_in_tree(db, person_id, promote_children=False)
            flash("Data tarombo berhasil diperbarui.", "success")
        return redirect(url_for("admin_edit_person", person_id=person_id))

    sundut_entries = list_sundut_entries(db)
    siblings = get_siblings_for_person(db, person_id, sql_approved, get_person_by_name)
    return render_template(
        "admin_edit_person.html",
        person=person,
        children=get_children_of_parent(db, person_id),
        submitted_children=get_submitted_children_of_parent(db, person_id),
        siblings=siblings,
        sundut_entries=sundut_entries,
        partuturan_options=PARTUTURAN_STATUSES,
    )


@app.route("/admin/sundut", methods=["GET", "POST"])
@login_required
@staff_required
def admin_sundut():
    db = get_db()
    if request.method == "POST":
        action = request.form.get("action", "")
        if action == "add_entry":
            number = request.form.get("sundut_number", "").strip()
            title = request.form.get("title", "").strip()
            description = request.form.get("description", "").strip()
            if number.isdigit() and title:
                try:
                    db.execute(
                        """
                        INSERT INTO sundut_entries (sundut_number, title, description)
                        VALUES (?, ?, ?)
                        """,
                        (int(number), title, description or None),
                    )
                    db.commit()
                    flash("Entri sundut berhasil ditambahkan.", "success")
                except Exception:
                    db.rollback()
                    flash("Nomor sundut sudah ada atau tidak valid.", "error")
            else:
                flash("Nomor sundut dan judul wajib diisi.", "error")
        elif action == "update_entry":
            entry_id = request.form.get("entry_id")
            title = request.form.get("title", "").strip()
            description = request.form.get("description", "").strip()
            db.execute(
                "UPDATE sundut_entries SET title = ?, description = ? WHERE id = ?",
                (title, description or None, entry_id),
            )
            db.commit()
            flash("Entri sundut diperbarui.", "success")
        elif action == "delete_entry":
            entry_id = request.form.get("entry_id")
            db.execute("DELETE FROM sundut_entries WHERE id = ?", (entry_id,))
            db.commit()
            flash("Entri sundut dihapus.", "success")
        elif action == "sync_sundut":
            result = sync_tarombo(
                db,
                sql_approved,
                insert_relationship_ignore,
                backfill=False,
                merge=False,
                purge=False,
            )
            log_audit(db, session.get("user_id"), "graph", None, "sync_sundut", result.get("sundut"))
            db.commit()
            sundut = result.get("sundut") or {}
            flash(sundut.get("message", "Sinkron sundut selesai."), "success")
        return redirect(url_for("admin_sundut"))

    entries = list_sundut_entries(db)
    assigned = db.fetchall(
        f"""
        SELECT sundut, COUNT(*) AS c
        FROM people
        WHERE {sql_approved(db)} AND sundut IS NOT NULL
        GROUP BY sundut
        ORDER BY sundut ASC
        """
    )
    assigned_map = {int(r["sundut"]): int(r["c"]) for r in assigned if r.get("sundut") is not None}
    return render_template(
        "admin_sundut.html",
        entries=entries,
        assigned_map=assigned_map,
    )


@app.route("/admin", methods=["GET", "POST"])
@login_required
@staff_required
def admin():
    db = get_db()
    if request.method == "POST":
        person_id = request.form.get("person_id")
        user_id = request.form.get("user_id")
        action = request.form.get("action")
        if action == "approve_member":
            if not user_id:
                flash("Member tidak ditemukan.", "error")
            else:
                approve_member_registration(db, user_id, actor_user_id=session.get("user_id"))
                db.commit()
                flash("Pendaftaran disetujui. Member dapat login dan melihat Tarombo.", "success")
        elif action == "reject_member":
            if not user_id:
                flash("Member tidak ditemukan.", "error")
            else:
                reject_member_registration(db, user_id, actor_user_id=session.get("user_id"))
                db.commit()
                flash("Pendaftaran ditolak.", "error")
        elif action == "approve":
            row = db.fetchone("SELECT submitted_by FROM people WHERE id = ?", (person_id,))
            if row and row.get("submitted_by"):
                approve_member_registration(db, row["submitted_by"], actor_user_id=session.get("user_id"))
            elif person_id:
                db.execute("UPDATE people SET approved = ? WHERE id = ?", (bool_db(db, True), person_id))
                place_person_in_tree(db, int(person_id))
                log_audit(db, session.get("user_id"), "person", person_id, "approve", {"approved": True})
            db.commit()
            flash("Data tarombo disetujui, tampil di pohon publik, dan member diaktifkan.", "success")
        elif action == "reject":
            db.execute(
                "UPDATE users SET member_status = 'rejected' WHERE id = (SELECT submitted_by FROM people WHERE id = ?)",
                (person_id,),
            )
            log_audit(db, session.get("user_id"), "person", person_id, "reject", {})
            db.commit()
            flash("Pendaftaran ditolak.", "error")
        elif action == "resync_all":
            resync_graph(db)
            log_audit(db, session.get("user_id"), "graph", None, "resync_all", {})
            db.commit()
            flash(
                "Sinkron lengkap: relasi ayah, panggoaran, sundut (+1 per generasi dari nilai manual).",
                "success",
            )
        elif action in ("merge_duplicates", "purge_non_tree"):
            result = reconcile_tree(db, sql_approved, insert_relationship_ignore)
            log_audit(db, session.get("user_id"), "graph", None, action, result)
            db.commit()
            merged = result["merge"]["people_merged"]
            deleted = result["purge"]["deleted"]
            if merged or deleted:
                parts = []
                if merged:
                    parts.append(
                        f"{merged} entri digabung (nama + ayah + ibu sama)"
                    )
                if deleted:
                    parts.append(f"{deleted} entri di luar pohon dihapus")
                flash(
                    "Pembersihan selesai: " + "; ".join(parts) + f". "
                    f"Pohon kini {result['tree_people']} orang.",
                    "success",
                )
            else:
                flash(
                    f"Tidak ada duplikat atau entri luar pohon. Pohon: {result['tree_people']} orang.",
                    "success",
                )
        elif action == "delete_person":
            db.execute("DELETE FROM relationships WHERE source_person_id = ? OR target_person_id = ?", (person_id, person_id))
            db.execute("DELETE FROM people WHERE id = ?", (person_id,))
            log_audit(db, session.get("user_id"), "person", person_id, "delete", {})
            db.commit()
            flash("Data tarombo berhasil dihapus.", "success")
        elif action == "delete_member":
            db.execute(
                """
                DELETE FROM relationships WHERE source_person_id IN (
                  SELECT id FROM people WHERE submitted_by = ?
                ) OR target_person_id IN (SELECT id FROM people WHERE submitted_by = ?)
                """,
                (user_id, user_id),
            )
            db.execute("DELETE FROM people WHERE submitted_by = ?", (user_id,))
            db.execute("DELETE FROM users WHERE id = ? AND role = 'member'", (user_id,))
            log_audit(db, session.get("user_id"), "user", user_id, "delete_member", {})
            db.commit()
            flash("Member dan data tarombo terkait berhasil dihapus.", "success")
        return redirect(url_for("admin"))

    pending_members = fetch_pending_registrations(db)
    pending = db.fetchall(
        f"""
        SELECT p.id, p.full_name, p.gender, p.marga, p.tarombo_status, p.father_name, p.mother_name,
               p.reference_female_line_name, p.spouse_name, p.spouse_marga,
               u.full_name AS submitter_name,
               (SELECT COUNT(*) FROM submitted_children sc WHERE sc.parent_person_id = p.id) AS child_count
        FROM people p
        JOIN users u ON u.id = p.submitted_by
        WHERE {sql_not_approved(db, 'p')}
        ORDER BY p.created_at DESC
        """
    )
    member_query = (request.args.get("q") or "").strip()
    if member_query:
        like = f"%{member_query.lower()}%"
        members = db.fetchall(
            """
            SELECT id, full_name, email, member_status, created_at FROM users
            WHERE role = 'member' AND (LOWER(full_name) LIKE ? OR LOWER(email) LIKE ?)
            ORDER BY created_at DESC
            """,
            (like, like),
        )
    else:
        members = db.fetchall(
            "SELECT id, full_name, email, member_status, created_at FROM users WHERE role = 'member' ORDER BY created_at DESC"
        )
    roots = build_collapsible_tree(db, sync=False)
    tree_ids = collect_tree_person_ids(roots)
    if tree_ids:
        placeholders = ",".join("?" * len(tree_ids))
        approved_people = db.fetchall(
            f"""
            SELECT p.id, p.full_name, p.gender, p.marga, p.tarombo_status, p.father_name, p.mother_name,
                   p.panggoaran, p.panggoaran_type, p.sundut, p.created_at, u.full_name AS submitter_name
            FROM people p
            LEFT JOIN users u ON u.id = p.submitted_by
            WHERE {sql_approved(db, 'p')} AND p.id IN ({placeholders})
            ORDER BY COALESCE(p.sundut, 9999) ASC, p.full_name ASC
            """,
            tuple(tree_ids),
        )
    else:
        approved_people = []
    duplicate_people = find_duplicate_identity_groups(db)
    off_tree_count = db.fetchone(
        f"""
        SELECT COUNT(*) AS c FROM people
        WHERE {sql_approved(db)} AND full_name != 'Tidak Diketahui'
        """
    )
    off_tree = max(0, int(off_tree_count["c"] or 0) - len(tree_ids))
    return render_template(
        "admin.html",
        pending_members=pending_members,
        pending=pending,
        members=members,
        member_query=member_query,
        approved_people=approved_people,
        duplicate_people=duplicate_people,
        tree_people_count=len(tree_ids),
        off_tree_count=off_tree,
    )


MEMBER_CSV_COLUMNS = ["full_name", "email", "password", "role", "member_status"]
_IMPORT_ALLOWED_ROLES = {"member", "pengurus"}
_IMPORT_ALLOWED_STATUS = {"pending", "active", "rejected"}


def _normalize_csv_header(name):
    return (name or "").strip().lower().replace(" ", "_").lstrip("\ufeff")


def _parse_members_csv(file_storage):
    """Baca CSV anggota → daftar dict {full_name,email,password,role,member_status}."""
    raw = file_storage.read()
    if isinstance(raw, bytes):
        text = raw.decode("utf-8-sig", errors="replace")
    else:
        text = str(raw)
    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    if not reader.fieldnames:
        return []
    field_map = {_normalize_csv_header(f): f for f in reader.fieldnames}
    rows = []
    for raw_row in reader:
        def col(key):
            src = field_map.get(key)
            return (raw_row.get(src) or "").strip() if src else ""

        full_name = col("full_name") or col("nama") or col("name")
        email = (col("email") or col("surel")).lower()
        if not full_name and not email:
            continue
        role = (col("role") or col("peran")).lower()
        if role not in _IMPORT_ALLOWED_ROLES:
            role = "member"
        status = (col("member_status") or col("status")).lower()
        if status not in _IMPORT_ALLOWED_STATUS:
            status = "active"
        rows.append(
            {
                "full_name": full_name,
                "email": email,
                "password": col("password") or col("kata_sandi") or "12345678",
                "role": role,
                "member_status": status,
            }
        )
    return rows


@app.route("/admin/import-members", methods=["GET", "POST"])
@login_required
@staff_required
def admin_import_members():
    """Impor akun anggota dari CSV (email + nama → buat login)."""
    db = get_db()
    summary = None
    if request.method == "POST":
        upload = request.files.get("csv_file")
        if not upload or not upload.filename:
            flash("Pilih berkas CSV terlebih dahulu.", "error")
            return redirect(url_for("admin_import_members"))
        try:
            rows = _parse_members_csv(upload)
        except Exception as exc:
            app.logger.exception("parse members csv")
            flash(f"Gagal membaca CSV: {exc}", "error")
            return redirect(url_for("admin_import_members"))

        created, skipped, errors = 0, 0, []
        for idx, row in enumerate(rows, start=1):
            email = row["email"]
            full_name = row["full_name"]
            if not email or "@" not in email:
                errors.append(f"Baris {idx}: email tidak valid ('{email}').")
                continue
            if not full_name:
                errors.append(f"Baris {idx}: nama kosong untuk {email}.")
                continue
            if db.fetchone("SELECT id FROM users WHERE email = ?", (email,)):
                skipped += 1
                continue
            password = row["password"] if len(row["password"]) >= 6 else "12345678"
            try:
                insert_returning_id(
                    db,
                    """
                    INSERT INTO users (full_name, email, password_hash, role, member_status)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (full_name, email, generate_password_hash(password), row["role"], row["member_status"]),
                )
                created += 1
            except Exception as exc:
                db.rollback()
                errors.append(f"Baris {idx}: gagal menyimpan {email} ({exc}).")
        db.commit()
        log_audit(
            db,
            session.get("user_id"),
            "users",
            None,
            "import_members_csv",
            {"created": created, "skipped": skipped, "errors": len(errors)},
        )
        db.commit()
        summary = {"created": created, "skipped": skipped, "errors": errors, "total": len(rows)}
        flash(
            f"Impor selesai: {created} dibuat, {skipped} dilewati (email sudah ada), "
            f"{len(errors)} error.",
            "success" if created else "error",
        )

    return render_template(
        "admin_import_members.html",
        summary=summary,
        columns=MEMBER_CSV_COLUMNS,
    )


@app.route("/admin/import-members/template.csv")
@login_required
@staff_required
def admin_import_members_template():
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(MEMBER_CSV_COLUMNS)
    writer.writerow(["Budi Samosir", "budi@example.com", "12345678", "member", "active"])
    writer.writerow(["Citra Pengurus", "citra@example.com", "12345678", "pengurus", "active"])
    data = buffer.getvalue().encode("utf-8-sig")
    return send_file(
        io.BytesIO(data),
        mimetype="text/csv",
        as_attachment=True,
        download_name="template-anggota.csv",
    )


@app.route("/admin/export-members")
@login_required
@staff_required
def admin_export_members():
    """Unduh daftar anggota sebagai XLSX (fallback CSV bila openpyxl tak ada)."""
    db = get_db()
    members = db.fetchall(
        """
        SELECT id, full_name, email, role, member_status, created_at
        FROM users
        ORDER BY role ASC, full_name ASC
        """
    )
    headers = ["ID", "Nama", "Email", "Peran", "Status", "Terdaftar"]
    stamp = datetime.now().strftime("%Y%m%d-%H%M")

    def row_values(m):
        return [
            m["id"],
            m["full_name"],
            m["email"],
            role_label(m["role"]),
            m["member_status"],
            str(m["created_at"]),
        ]

    try:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Anggota"
        ws.append(headers)
        for m in members:
            ws.append(row_values(m))
        for i, _ in enumerate(headers, start=1):
            ws.column_dimensions[chr(64 + i)].width = 22
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        log_audit(db, session.get("user_id"), "users", None, "export_members_xlsx", {"count": len(members)})
        db.commit()
        return send_file(
            buffer,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"anggota-ptsbi-{stamp}.xlsx",
        )
    except ImportError:
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(headers)
        for m in members:
            writer.writerow(row_values(m))
        data = buffer.getvalue().encode("utf-8-sig")
        log_audit(db, session.get("user_id"), "users", None, "export_members_csv", {"count": len(members)})
        db.commit()
        flash("openpyxl belum terpasang; mengunduh format CSV sebagai gantinya.", "success")
        return send_file(
            io.BytesIO(data),
            mimetype="text/csv",
            as_attachment=True,
            download_name=f"anggota-ptsbi-{stamp}.csv",
        )


@app.route("/admin/form-config", methods=["GET", "POST"])
@login_required
@developer_required
def admin_form_config():
    db = get_db()
    if request.method == "POST":
        fields, _ = get_all_fields_admin(db)
        for field in fields:
            fid = field["id"]
            db.execute(
                """
                UPDATE form_field_configs SET
                  label = ?, field_type = ?, display_order = ?,
                  is_required = ?, is_visible = ?, default_value = ?, help_text = ?
                WHERE id = ?
                """,
                (
                    request.form.get(f"label_{fid}", field["label"]).strip(),
                    request.form.get(f"type_{fid}", field["field_type"]).strip(),
                    int(request.form.get(f"order_{fid}", field["display_order"]) or 0),
                    bool_db(db, f"req_{fid}" in request.form),
                    bool_db(db, f"vis_{fid}" in request.form),
                    request.form.get(f"default_{fid}", "").strip() or None,
                    request.form.get(f"help_{fid}", "").strip() or None,
                    fid,
                ),
            )
        log_audit(db, session.get("user_id"), "form_config", None, "update", {"fields": len(fields)})
        db.commit()
        flash("Pengaturan form berhasil disimpan.", "success")
        return redirect(url_for("admin_form_config"))

    fields, section_titles = get_all_fields_admin(db)
    return render_template("admin_form_config.html", fields=fields, section_titles=section_titles)


@app.route("/admin/register-member", methods=["GET", "POST"])
@app.route("/admin/tambah-anggota", methods=["GET", "POST"])
@login_required
@staff_required
def admin_add_tarombo_person():
    """Admin menambah anggota ke pohon saja — tanpa membuat akun user."""
    db = get_db()
    ctx = admin_tarombo_form_context(db)
    if request.method == "POST":
        try:
            result = add_tarombo_person_from_request(
                db,
                request.form,
                admin_user_id=session.get("user_id"),
            )
        except ValueError as exc:
            flash(str(exc), "error")
            return render_template("admin_add_person.html", **ctx)
        log_audit(
            db,
            session.get("user_id"),
            "person",
            result["person_id"],
            "admin_add_tarombo",
            {"full_name": result["full_name"]},
        )
        db.commit()
        flash(
            f"{result['full_name']} ditambahkan ke tarombo dan sudah tampil di pohon.",
            "success",
        )
        return redirect(url_for("admin_edit_person", person_id=result["person_id"]))

    return render_template("admin_add_person.html", **ctx)


def _run_database_reset():
    global _db_bootstrapped
    db = get_db()
    backend = db.backend
    if backend == "postgres":
        wipe_all_data(db)
    reset_request_db()
    if backend == "sqlite":
        reset_database_file_sqlite()
    _db_bootstrapped = False
    db = get_db()
    if backend == "sqlite":
        init_schema(db)
    migrate_schema(db)
    backfill_normalized_names(db)
    seed_form_fields(db)
    seed_sundut_entries(db)
    seed_admin_only(db)
    _db_bootstrapped = True


@app.route("/admin/database", methods=["GET", "POST"])
@login_required
@developer_required
def admin_database():
    db = get_db()
    if request.method == "POST":
        action = request.form.get("action", "")
        try:
            if action == "backup":
                info = create_backup(db)
                log_audit(db, session.get("user_id"), "database", None, "backup", info)
                db.commit()
                flash(f"Backup berhasil: {info['filename']}", "success")
            elif action == "restore":
                filename = request.form.get("filename", "").strip()
                user_id = session.get("user_id")
                restore_backup(filename, close_db_fn=close_database_for_restore)
                bootstrap_after_db_restore()
                db = get_db()
                log_audit(
                    db,
                    user_id,
                    "database",
                    None,
                    "restore",
                    {"filename": filename},
                )
                db.commit()
                flash(f"Database dipulihkan dari {filename}. Silakan login ulang jika perlu.", "success")
            elif action == "restore_upload":
                upload = request.files.get("backup_file")
                if not upload or not upload.filename:
                    raise ValueError("Pilih file backup (.db atau .sql).")
                user_id = session.get("user_id")
                tmp_dir = BACKUP_DIR / "_upload_tmp"
                tmp_dir.mkdir(parents=True, exist_ok=True)
                safe_name = secure_filename(upload.filename)
                tmp_path = tmp_dir / safe_name
                upload.save(tmp_path)
                try:
                    restore_uploaded_file(
                        tmp_path, safe_name, close_db_fn=close_database_for_restore
                    )
                    bootstrap_after_db_restore()
                    db = get_db()
                    log_audit(
                        db,
                        user_id,
                        "database",
                        None,
                        "restore_upload",
                        {"filename": safe_name},
                    )
                    db.commit()
                    flash("Database dipulihkan dari file upload.", "success")
                finally:
                    try:
                        tmp_path.unlink(missing_ok=True)
                    except OSError:
                        pass
            elif action == "reset":
                confirm = request.form.get("confirm_text", "").strip()
                if confirm != "RESET":
                    raise ValueError('Ketik RESET (huruf besar) untuk mengonfirmasi reset database.')
                _run_database_reset()
                db = get_db()
                log_audit(db, session.get("user_id"), "database", None, "reset", {})
                db.commit()
                flash(
                    "Database direset. Akun bawaan: admin/pengurus/anggota/developer@ptsbi.org "
                    "(password 12345678). Login ulang.",
                    "success",
                )
                session.clear()
                return redirect(url_for("login"))
            else:
                flash("Aksi tidak dikenali.", "error")
        except Exception as exc:
            app.logger.exception("admin_database")
            flash(str(exc), "error")
        return redirect(url_for("admin_database"))

    backups = list_backups()
    return render_template(
        "admin_database.html",
        backups=backups,
        db_backend=db.backend,
        backup_dir=str(BACKUP_DIR),
    )


@app.route("/admin/database/download/<filename>")
@login_required
@developer_required
def admin_database_download(filename):
    path = resolve_backup_path(filename)
    return send_file(path, as_attachment=True, download_name=path.name)


@app.route("/admin/logo", methods=["GET", "POST"])
@login_required
@developer_required
def admin_logo():
    db = get_db()
    ensure_upload_dir()
    if request.method == "POST":
        action = request.form.get("action", "")
        if action == "upload":
            file = request.files.get("logo_file")
            try:
                save_logo_upload(file)
                if db.backend == "postgres":
                    db.execute(
                        """
                        INSERT INTO app_settings (key, value, updated_at)
                        VALUES ('logo', 'custom', NOW())
                        ON CONFLICT (key) DO UPDATE SET value = 'custom', updated_at = NOW()
                        """
                    )
                else:
                    db.execute(
                        """
                        INSERT OR REPLACE INTO app_settings (key, value, updated_at)
                        VALUES ('logo', 'custom', CURRENT_TIMESTAMP)
                        """
                    )
                log_audit(db, session.get("user_id"), "site", None, "logo_upload", {"file": "site-logo.png"})
                db.commit()
                flash("Logo berhasil diunggah dan ditampilkan di seluruh situs.", "success")
            except ValueError as exc:
                flash(str(exc), "error")
        elif action == "reset":
            if remove_custom_logo():
                db.execute("DELETE FROM app_settings WHERE key = 'logo'")
                log_audit(db, session.get("user_id"), "site", None, "logo_reset", {})
                db.commit()
                flash("Logo dikembalikan ke logo bawaan.", "success")
            else:
                flash("Tidak ada logo kustom untuk dihapus.", "success")
        return redirect(url_for("admin_logo"))
    return render_template("admin_logo.html", logo=logo_info())


@app.route("/admin/audit")
@login_required
@developer_required
def admin_audit():
    db = get_db()
    logs = db.fetchall(
        """
        SELECT a.*, u.full_name AS actor_name
        FROM audit_logs a
        LEFT JOIN users u ON u.id = a.actor_user_id
        ORDER BY a.created_at DESC
        LIMIT 200
        """
    )
    return render_template("admin_audit.html", logs=logs)


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5000"))
    debug = os.environ.get("FLASK_DEBUG", "1") == "1"
    app.run(host="0.0.0.0", port=port, debug=debug, use_reloader=debug)
